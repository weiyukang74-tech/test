from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from evidence_retriever.planning_support import collect_progress_records


class CollectProgressRecordsTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        self.mapping_path = root / "mapping.xlsx"
        self.data_dir = root / "data"
        self.data_dir.mkdir()
        self._write_mapping()
        self._write_medical_data()

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _write_mapping(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "拆分映射"
        worksheet.append(
            [
                "mapping_id",
                "medinsId",
                "mdtrtId",
                "来源级别",
                "主要来源",
                "次要来源",
                "对应的视图",
                "文书类型",
                "一级标签名",
                "二级标签名",
                "匹配到的recordName",
                "recordType",
                "匹配方式",
                "匹配状态",
            ]
        )
        for mapping_id, source, record_name, record_type, match_mode in (
            ("M01", "DS-M-IP-004 日常病程记录", "日常病程记录", "病程记录", "recordName精确匹配"),
            ("M02", "DS-M-IP-005 上级医师查房记录", "查房记录", "病程记录", "recordName包含"),
            ("M03", "DS-M-IP-002 入院记录", None, "入院记录", "recordType精确匹配"),
        ):
            worksheet.append(
                [
                    mapping_id,
                    "运行时传入",
                    "运行时传入",
                    "PRIMARY",
                    source,
                    "—",
                    "病历信息",
                    "病历",
                    "治疗",
                    "治疗过程",
                    record_name,
                    record_type,
                    match_mode,
                    "可直接查询",
                ]
            )
        worksheet.append(
            [
                "M04",
                "运行时传入",
                "运行时传入",
                "PRIMARY",
                "DS-S-001 诊断信息",
                "—",
                "诊断信息",
                "具体信息",
                "诊断",
                "标准诊断",
                None,
                None,
                "按mdtrtId+medinsId查询",
                "可直接查询",
            ]
        )
        workbook.save(self.mapping_path)
        workbook.close()

    def _write_medical_data(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "病历信息"
        worksheet.append(
            ["medinsId", "mdtrtId", "recordName", "recordType", "updateTime", "recordContent"]
        )
        for index, (record_name, record_type) in enumerate(
            (
                ("日常病程记录", "病程记录"),
                ("张三主任医师查房记录", "病程记录"),
                ("入院记录", "入院记录"),
            ),
            start=1,
        ):
            worksheet.append(
                [
                    "HOSP-1",
                    "CASE-1",
                    record_name,
                    record_type,
                    f"2026-01-0{index} 08:00:00",
                    f"{record_name}：继续床边康复治疗。",
                ]
            )
        workbook.save(self.data_dir / "病历信息.xlsx")
        workbook.close()

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "患者信息"
        worksheet.append(["medinsId", "mdtrtId", "patnName", "gend"])
        worksheet.append(["HOSP-1", "CASE-1", "测试患者", "2"])
        workbook.save(self.data_dir / "患者信息.xlsx")
        workbook.close()

    def test_returns_only_progress_records_for_mapped_tags(self) -> None:
        collected = collect_progress_records(
            self.mapping_path,
            self.data_dir,
            "HOSP-1",
            "CASE-1",
            [("治疗", "治疗过程"), ("诊断", "标准诊断")],
        )

        self.assertEqual(
            [record["recordName"] for record in collected[("治疗", "治疗过程")]],
            ["日常病程记录", "张三主任医师查房记录"],
        )
        self.assertEqual(collected[("诊断", "标准诊断")], [])

    def test_resolves_medins_id_placeholder_from_data(self) -> None:
        collected = collect_progress_records(
            self.mapping_path,
            self.data_dir,
            "运行时传入",
            "CASE-1",
            [("治疗", "治疗过程")],
        )

        self.assertEqual(len(collected[("治疗", "治疗过程")]), 2)

    def test_unknown_tag_returns_empty_list(self) -> None:
        collected = collect_progress_records(
            self.mapping_path,
            self.data_dir,
            "HOSP-1",
            "CASE-1",
            [("护理", "护理记录")],
        )

        self.assertEqual(collected[("护理", "护理记录")], [])


if __name__ == "__main__":
    unittest.main()
