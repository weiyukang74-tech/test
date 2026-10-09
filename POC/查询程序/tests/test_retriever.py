from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from evidence_retriever.context_builder import MedicalDocumentCollector, build_review_context
from evidence_retriever.service import RetrievalService


PROGRAM_ROOT = Path(__file__).resolve().parents[1]
POC_ROOT = PROGRAM_ROOT.parent
WORKSPACE_ROOT = POC_ROOT.parent
RULE_EVIDENCE = POC_ROOT / "规则证据结果" / "规则证据结果.json"
MAPPING = POC_ROOT / "标签视图映射" / "规则证据标签关系映射表_v2.xlsx"
MEDICAL_DATA = WORKSPACE_ROOT / "datat" / "脱敏病历.md"


class RetrievalServiceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.service = RetrievalService(MAPPING, MEDICAL_DATA)
        self.rules = self.service.load_rule_evidence(RULE_EVIDENCE)

    def test_p01_has_compact_tag_results_and_full_medical_record(self) -> None:
        result = self.service.run(
            self.rules,
            medins_id="H44030500096",
            mdtrt_id="119797913",
            proposition_id="P01",
            secondary_mode="auto",
        )

        self.assertEqual(set(result), {"rule_id", "propositions"})
        self.assertEqual(result["rule_id"], "RULE-001")
        self.assertEqual(len(result["propositions"]), 1)
        proposition = result["propositions"][0]
        self.assertEqual(proposition["proposition_id"], "P01")

        by_tag = {
            (entry["一级证据域"], entry["二级证据标签"]): entry
            for entry in proposition["evidence"]
        }
        self.assertEqual(
            by_tag[("患者", "性别")]["query_result"]["data"],
            [{"gend": "2"}],
        )

        reproductive = by_tag[("患者", "生育状态")]["query_result"]
        self.assertEqual(reproductive["status"], "RETRIEVED")
        self.assertEqual(reproductive["source"], "DS-M-IP-002 入院记录")
        record = reproductive["data"][0]
        self.assertEqual(record["recordName"], "再次入院记录")
        self.assertGreater(len(record["recordContent"]), 20)
        self.assertIn("现病史", record["recordContent"])
        self.assertIn("绝经后", record["recordContent"]["月经史"])

        pathology = by_tag[("病理", "病理标志")]["query_result"]
        self.assertEqual(pathology["status"], "NOT_RETRIEVED")
        self.assertEqual(pathology["data"], [])

    def test_diagnosis_keyword_comes_from_target_fact(self) -> None:
        result = self.service.run(
            self.rules,
            medins_id="H44030500096",
            mdtrt_id="119797913",
            proposition_id="P01",
            secondary_mode="auto",
        )
        evidence = result["propositions"][0]["evidence"]
        diagnosis = next(
            entry for entry in evidence if entry["二级证据标签"] == "标准诊断"
        )["query_result"]

        self.assertEqual(
            diagnosis["matched_by"],
            {
                "field": "diseName",
                "operator": "contains",
                "value": "乳腺癌",
                "value_source": "target_fact",
            },
        )
        self.assertEqual(len(diagnosis["data"]), 2)
        self.assertTrue(
            all("乳腺癌" in record["diseName"] for record in diagnosis["data"])
        )

    def test_full_output_is_json_serializable(self) -> None:
        result = self.service.run(
            self.rules,
            medins_id="H44030500096",
            mdtrt_id="119797913",
            secondary_mode="auto",
        )
        serialized = json.dumps(result, ensure_ascii=False)
        self.assertIn('"propositions"', serialized)
        self.assertNotIn('"evidence_pool"', serialized)
        self.assertNotIn('"summary"', serialized)

    def test_review_context_merges_rules_and_deduplicates_documents(self) -> None:
        query_result = self.service.run(
            self.rules,
            medins_id="H44030500096",
            mdtrt_id="119797913",
            secondary_mode="auto",
        )
        context = build_review_context(self.rules, query_result)

        self.assertEqual(
            set(context),
            {
                "rule_id",
                "review_basis",
                "propositions",
                "structured_data_pool",
                "medical_documents",
                "program_checks",
            },
        )
        self.assertEqual(context["rule_id"], "RULE-001")
        self.assertEqual(context["review_basis"]["violation_item"], "依西美坦片")
        self.assertEqual(len(context["propositions"]), 3)

        p01 = next(
            proposition
            for proposition in context["propositions"]
            if proposition["proposition_id"] == "P01"
        )
        reproductive = next(
            item
            for item in p01["evidence"]
            if item["二级证据标签"] == "生育状态"
        )
        self.assertEqual(reproductive["necessity"], "REQUIRED")
        document_id = reproductive["query_result"]["document_refs"][0]

        document = next(
            item
            for item in context["medical_documents"]
            if item["document_id"] == document_id
        )
        self.assertEqual(document["recordName"], "再次入院记录")
        self.assertIn("现病史", document["recordContent"])
        self.assertIn("绝经后", document["recordContent"]["月经史"])

        all_document_refs = [
            document_ref
            for proposition in context["propositions"]
            for evidence in proposition["evidence"]
            for document_ref in evidence["query_result"].get("document_refs", [])
        ]
        self.assertGreater(len(all_document_refs), len(set(all_document_refs)))
        self.assertEqual(
            len(context["medical_documents"]),
            len(set(all_document_refs)),
        )

        diagnosis = next(
            item
            for item in p01["evidence"]
            if item["二级证据标签"] == "标准诊断"
        )
        self.assertNotIn("matched_by", diagnosis["query_result"])
        structured_refs = diagnosis["query_result"]["structured_data_refs"]
        self.assertEqual(
            len(structured_refs),
            2,
        )
        self.assertTrue(
            all(reference in context["structured_data_pool"] for reference in structured_refs)
        )


class MedicalDocumentCollectorTests(unittest.TestCase):
    def test_deduplicates_transport_variants_and_merges_matches(self) -> None:
        collector = MedicalDocumentCollector()
        first_record = {
            "medinsId": "HOSP-1",
            "mdtrtId": "CASE-1",
            "recordName": "日常病程记录",
            "recordType": "病程记录",
            "updateTime": "2026-01-01 08:00:00",
            "recordContent": "ER+，既往使用他莫昔芬。",
            "transport_page": 1,
            "_medical_record_filter_match": {
                "matched_propositions": ["P03"],
                "matched_evidence_tags": [{"proposition_id": "P03"}],
                "matched_terms": ["ER"],
            },
        }
        second_record = json.loads(json.dumps(first_record, ensure_ascii=False))
        second_record["transport_page"] = 9
        second_record["_medical_record_filter_match"] = {
            "matched_propositions": ["P04"],
            "matched_evidence_tags": [{"proposition_id": "P04"}],
            "matched_terms": ["他莫昔芬"],
        }

        first_id = collector.add("DS-M-IP-003 首次病程记录", first_record)
        second_id = collector.add("DS-M-IP-004 日常病程记录", second_record)

        self.assertEqual(first_id, second_id)
        self.assertEqual(len(collector.documents), 1)
        document = collector.documents[0]
        self.assertEqual(
            document["source"],
            "DS-M-IP-003 首次病程记录、DS-M-IP-004 日常病程记录",
        )
        self.assertEqual(document["matched_propositions"], ["P03", "P04"])
        self.assertEqual(
            document["matched_evidence_tags"],
            [{"proposition_id": "P03"}, {"proposition_id": "P04"}],
        )
        self.assertEqual(document["matched_terms"], ["ER", "他莫昔芬"])


class MedicalRecordFilteringIntegrationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        self.mapping_path = root / "mapping.xlsx"
        self.data_dir = root / "data"
        self.data_dir.mkdir()
        self._write_mapping()
        self._write_medical_data()
        self.service = RetrievalService(self.mapping_path, self.data_dir)
        self.rule = {
            "rule_id": "RULE-056",
            "result": {
                "source": {
                    "violation_item": "依西美坦片",
                    "violation_description": "患者开具此药属于超适应症用药",
                },
                "regulatory_claim": {"statement": "疑似超适应症"},
                "appeal_objective": "证明患者符合依西美坦适应症",
                "appeal_propositions": [
                    {
                        "proposition_id": "P03",
                        "statement": "患者雌激素受体阳性",
                        "logic_notes": "核实受体状态",
                        "required_tags": [
                            {
                                "一级证据域": "病理",
                                "二级证据标签": "病理结果",
                                "target_fact": "确认患者雌激素受体状态",
                                "necessity": "REQUIRED",
                            }
                        ],
                        "supplementary_tags": [],
                        "missing_required_facts": [],
                    },
                    {
                        "proposition_id": "P04",
                        "statement": "患者既往使用过他莫昔芬",
                        "logic_notes": "核实既往治疗",
                        "required_tags": [
                            {
                                "一级证据域": "用药",
                                "二级证据标签": "既往治疗",
                                "target_fact": "确认患者既往使用他莫昔芬",
                                "necessity": "REQUIRED",
                            }
                        ],
                        "supplementary_tags": [],
                        "missing_required_facts": [],
                    },
                ],
                "program_checks": [],
            },
        }
        self.query_plan = {
            "rule_id": "RULE-056",
            "plans": [
                {
                    "query_id": "QP-P03-01",
                    "supports_facts": [
                        {
                            "proposition_id": "P03",
                            "一级证据域": "病理",
                            "二级证据标签": "病理结果",
                            "target_fact": "确认患者雌激素受体状态",
                            "necessity": "REQUIRED",
                        }
                    ],
                    "query_type": "MEDICAL_RECORD_KEYWORD_FILTER",
                    "record_search_terms": ["雌激素", "ER"],
                    "medical_record_sources": ["DS-M-IP-004 日常病程记录"],
                },
                {
                    "query_id": "QP-P04-01",
                    "supports_facts": [
                        {
                            "proposition_id": "P04",
                            "一级证据域": "用药",
                            "二级证据标签": "既往治疗",
                            "target_fact": "确认患者既往使用他莫昔芬",
                            "necessity": "REQUIRED",
                        }
                    ],
                    "query_type": "MEDICAL_RECORD_KEYWORD_FILTER",
                    "record_search_terms": ["他莫昔芬"],
                    "medical_record_sources": ["DS-M-IP-004 日常病程记录"],
                },
            ],
        }

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _write_mapping(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "拆分映射"
        worksheet.append(
            [
                "mapping_id",
                "medinsId",
                "mdtrtId",
                "一级标签名",
                "二级标签名",
                "来源级别",
                "主要来源",
                "次要来源",
                "对应的视图",
                "文书类型",
                "匹配到的recordName",
                "recordType",
                "匹配状态",
            ]
        )
        for mapping_id, primary_domain, secondary_tag in (
            ("M01", "病理", "病理结果"),
            ("M02", "用药", "既往治疗"),
        ):
            worksheet.append(
                [
                    mapping_id,
                    "运行时传入",
                    "运行时传入",
                    primary_domain,
                    secondary_tag,
                    "PRIMARY",
                    "DS-M-IP-004 日常病程记录",
                    "—",
                    "病历信息",
                    "病历",
                    "日常病程记录",
                    "病程记录",
                    "已匹配",
                ]
            )
        workbook.save(self.mapping_path)
        workbook.close()

    def _write_medical_data(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "病历信息"
        worksheet.append(
            [
                "medinsId",
                "mdtrtId",
                "recordName",
                "recordType",
                "updateTime",
                "recordContent",
            ]
        )
        worksheet.append(
            [
                "HOSP-1",
                "CASE-1",
                "日常病程记录",
                "病程记录",
                "2026-01-01 08:00:00",
                json.dumps({"病理结果": "ER+，HER2阴性"}, ensure_ascii=False),
            ]
        )
        worksheet.append(
            [
                "HOSP-1",
                "CASE-1",
                "日常病程记录",
                "病程记录",
                "2026-01-02 08:00:00",
                "患者既往接受他莫昔芬辅助治疗。",
            ]
        )
        worksheet.append(
            [
                "HOSP-1",
                "CASE-1",
                "日常病程记录",
                "病程记录",
                "2026-01-03 08:00:00",
                "仅记录HER2阳性，患者一般情况稳定。",
            ]
        )
        workbook.save(self.data_dir / "病历信息.xlsx")
        workbook.close()

    def test_run_filters_progress_records_with_rule_level_or(self) -> None:
        result = self.service.run(
            [self.rule],
            medins_id="HOSP-1",
            mdtrt_id="CASE-1",
            rule_id="RULE-056",
            query_plan=self.query_plan,
        )

        self.assertEqual(
            result["medical_record_filter"],
            {
                "enabled": True,
                "filter_applied": True,
                "candidate_document_count": 3,
                "selected_document_count": 2,
                "filter_fallback": False,
                "fallback_reason": None,
                "planned_evidence_count": 2,
                "matched_evidence_count": 2,
                "required_planned_evidence_count": 2,
                "required_matched_evidence_count": 2,
            },
        )
        for proposition in result["propositions"]:
            query_result = proposition["evidence"][0]["query_result"]
            self.assertEqual(query_result["source"], "DS-M-IP-004 日常病程记录")
            self.assertEqual(len(query_result["data"]), 2)
            self.assertTrue(
                all(
                    "_medical_record_filter_match" in record
                    for record in query_result["data"]
                )
            )

    def test_run_falls_back_when_no_terms_match(self) -> None:
        query_plan = copy_query_plan = json.loads(json.dumps(self.query_plan))
        copy_query_plan["plans"][0]["record_search_terms"] = ["孕激素受体"]
        copy_query_plan["plans"][1]["record_search_terms"] = ["芳香化酶"]

        result = self.service.run(
            [self.rule],
            medins_id="HOSP-1",
            mdtrt_id="CASE-1",
            rule_id="RULE-056",
            query_plan=query_plan,
        )

        stats = result["medical_record_filter"]
        self.assertTrue(stats["filter_fallback"])
        self.assertEqual(stats["fallback_reason"], "NO_KEYWORD_MATCH")
        self.assertEqual(len(result["propositions"][0]["evidence"][0]["query_result"]["data"]), 3)

    def test_rule_required_tags_override_stale_plan_necessity(self) -> None:
        query_plan = json.loads(json.dumps(self.query_plan))
        query_plan["plans"][1]["supports_facts"][0]["necessity"] = "SUPPLEMENTARY"
        query_plan["plans"][1]["record_search_terms"] = ["芳香化酶"]

        result = self.service.run(
            [self.rule],
            medins_id="HOSP-1",
            mdtrt_id="CASE-1",
            rule_id="RULE-056",
            query_plan=query_plan,
        )

        stats = result["medical_record_filter"]
        self.assertTrue(stats["filter_fallback"])
        self.assertEqual(stats["fallback_reason"], "INCOMPLETE_FILTER_COVERAGE")
        self.assertEqual(stats["required_planned_evidence_count"], 2)
        self.assertEqual(stats["required_matched_evidence_count"], 1)

    def test_unmatched_supplementary_evidence_keeps_its_progress_documents(self) -> None:
        rule = json.loads(json.dumps(self.rule))
        p04 = rule["result"]["appeal_propositions"][1]
        p04["supplementary_tags"] = p04["required_tags"]
        p04["required_tags"] = []
        p04["supplementary_tags"][0]["necessity"] = "SUPPLEMENTARY"
        query_plan = json.loads(json.dumps(self.query_plan))
        query_plan["plans"][1]["supports_facts"][0]["necessity"] = "SUPPLEMENTARY"
        query_plan["plans"][1]["record_search_terms"] = ["芳香化酶"]

        result = self.service.run(
            [rule],
            medins_id="HOSP-1",
            mdtrt_id="CASE-1",
            rule_id="RULE-056",
            query_plan=query_plan,
        )

        stats = result["medical_record_filter"]
        self.assertTrue(stats["filter_applied"])
        self.assertEqual(stats["selected_document_count"], 3)
        self.assertEqual(
            stats["supplementary_evidence_fallbacks"][0]["二级证据标签"],
            "既往治疗",
        )
        required_data = result["propositions"][0]["evidence"][0]["query_result"]["data"]
        supplementary_data = result["propositions"][1]["evidence"][0]["query_result"][
            "data"
        ]
        self.assertEqual(len(required_data), 1)
        self.assertEqual(len(supplementary_data), 3)

        context = build_review_context([rule], result)

        self.assertEqual(context["medical_document_groups"], [])
        self.assertEqual(len(context["medical_documents"]), 3)
        self.assertEqual(
            sum("matched_terms" in document for document in context["medical_documents"]),
            1,
        )

    def test_review_context_preserves_filtered_progress_documents(self) -> None:
        result = self.service.run(
            [self.rule],
            medins_id="HOSP-1",
            mdtrt_id="CASE-1",
            rule_id="RULE-056",
            query_plan=self.query_plan,
        )

        context = build_review_context([self.rule], result)

        self.assertEqual(context["medical_record_filter"]["selected_document_count"], 2)
        self.assertEqual(len(context["medical_documents"]), 2)
        self.assertEqual(context["medical_document_groups"], [])
        self.assertTrue(
            all(document["recordType"] == "病程记录" for document in context["medical_documents"])
        )
        self.assertTrue(
            all("matched_terms" in document for document in context["medical_documents"])
        )
        self.assertIn("ER+", json.dumps(context["medical_documents"], ensure_ascii=False))

    def test_missing_plan_keeps_legacy_document_grouping(self) -> None:
        result = self.service.run(
            [self.rule],
            medins_id="HOSP-1",
            mdtrt_id="CASE-1",
            rule_id="RULE-056",
            query_plan=None,
        )

        self.assertFalse(result["medical_record_filter"]["enabled"])
        context = build_review_context([self.rule], result)

        self.assertTrue(context["medical_document_groups"])

    def test_duplicate_medical_plan_safely_falls_back(self) -> None:
        query_plan = json.loads(json.dumps(self.query_plan))
        query_plan["plans"].append(json.loads(json.dumps(query_plan["plans"][0])))

        result = self.service.run(
            [self.rule],
            medins_id="HOSP-1",
            mdtrt_id="CASE-1",
            rule_id="RULE-056",
            query_plan=query_plan,
        )

        stats = result["medical_record_filter"]
        self.assertFalse(stats["enabled"])
        self.assertEqual(stats["fallback_reason"], "INCOMPLETE_FILTER_PLAN")
        self.assertTrue(stats["plan_validation_errors"])
        self.assertEqual(
            len(result["propositions"][0]["evidence"][0]["query_result"]["data"]),
            3,
        )


class RecordNameMatchModeTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        root = Path(self.temp_dir.name)
        mapping_path = root / "mapping.xlsx"
        data_dir = root / "data"
        data_dir.mkdir()

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "拆分映射"
        worksheet.append(
            [
                "mapping_id",
                "medinsId",
                "mdtrtId",
                "一级标签名",
                "二级标签名",
                "来源级别",
                "主要来源",
                "次要来源",
                "对应的视图",
                "文书类型",
                "匹配到的recordName",
                "recordType",
                "匹配方式",
                "匹配状态",
            ]
        )
        for mapping_id, secondary_tag, source, record_name, record_type, match_mode in (
            ("M01", "治疗过程", "DS-M-IP-004 日常病程记录", "日常病程记录", "病程记录", "recordName精确匹配"),
            ("M02", "治疗过程", "DS-M-IP-005 上级医师查房记录", "查房记录", "病程记录", "recordName包含"),
            ("M03", "治疗方案", "DS-M-IP-010 术前讨论记录", "术前病例讨论", "病程记录；手术相关", "recordName包含"),
        ):
            worksheet.append(
                [
                    mapping_id,
                    "运行时传入",
                    "运行时传入",
                    "治疗",
                    secondary_tag,
                    "PRIMARY",
                    source,
                    "—",
                    "病历信息",
                    "病历",
                    record_name,
                    record_type,
                    match_mode,
                    "可直接查询",
                ]
            )
        workbook.save(mapping_path)
        workbook.close()

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "病历信息"
        worksheet.append(
            ["medinsId", "mdtrtId", "recordName", "recordType", "updateTime", "recordContent"]
        )
        for index, (record_name, record_type) in enumerate(
            (
                ("日常病程记录", "病程记录"),
                ("日常病程记录（补）", "病程记录"),
                ("张三主任医师查房记录", "病程记录"),
                ("李四副主任医师代主治医师查房记录", "病程记录"),
                ("会诊记录", "病程记录"),
                ("主诊组术前病例讨论记录", "手术相关"),
                ("主诊组术前病例讨论结论记录", "病程记录"),
                ("手术记录", "手术相关"),
            ),
            start=1,
        ):
            worksheet.append(
                [
                    "HOSP-1",
                    "CASE-1",
                    record_name,
                    record_type,
                    f"2026-01-0{index} 08:00:00",
                    f"{record_name}：继续床边康复治疗。",
                ]
            )
        workbook.save(data_dir / "病历信息.xlsx")
        workbook.close()

        self.service = RetrievalService(mapping_path, data_dir)
        self.rule = {
            "rule_id": "RULE-009",
            "result": {
                "appeal_propositions": [
                    {
                        "proposition_id": "P01",
                        "required_tags": [
                            {
                                "一级证据域": "治疗",
                                "二级证据标签": "治疗过程",
                                "target_fact": "核对康复治疗执行情况",
                                "necessity": "REQUIRED",
                            }
                        ],
                        "supplementary_tags": [],
                    }
                ]
            },
        }

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_record_name_follows_mapping_match_mode(self) -> None:
        result = self.service.run(
            [self.rule],
            medins_id="HOSP-1",
            mdtrt_id="CASE-1",
            rule_id="RULE-009",
        )

        query_result = result["propositions"][0]["evidence"][0]["query_result"]
        self.assertEqual(
            [record["recordName"] for record in query_result["data"]],
            ["日常病程记录", "张三主任医师查房记录", "李四副主任医师代主治医师查房记录"],
        )
        self.assertEqual(
            query_result["source"],
            "DS-M-IP-004 日常病程记录、DS-M-IP-005 上级医师查房记录",
        )

    def test_record_type_with_semicolon_matches_any_listed_type(self) -> None:
        rule = json.loads(json.dumps(self.rule))
        tag = rule["result"]["appeal_propositions"][0]["required_tags"][0]
        tag["二级证据标签"] = "治疗方案"
        tag["target_fact"] = "查看术前讨论确定的治疗方案"

        result = self.service.run(
            [rule],
            medins_id="HOSP-1",
            mdtrt_id="CASE-1",
            rule_id="RULE-009",
        )

        query_result = result["propositions"][0]["evidence"][0]["query_result"]
        self.assertEqual(
            [(record["recordName"], record["recordType"]) for record in query_result["data"]],
            [
                ("主诊组术前病例讨论记录", "手术相关"),
                ("主诊组术前病例讨论结论记录", "病程记录"),
            ],
        )


if __name__ == "__main__":
    unittest.main()
