from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook


def load_expense_tags(path: Path) -> set[tuple[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["拆分映射"]
        header: dict[str, int] | None = None
        tags: set[tuple[str, str]] = set()
        for values in worksheet.iter_rows(values_only=True):
            if values and values[0] == "mapping_id":
                header = {
                    str(value): index
                    for index, value in enumerate(values)
                    if value is not None
                }
                continue
            if header is None or not values or not values[0]:
                continue
            view_name = str(values[header["对应的视图"]] or "").strip()
            source_cells = [
                str(values[header[name]] or "")
                for name in ("主要来源", "次要来源")
                if name in header
            ]
            if view_name != "费用明细" and not any(
                re.search(r"DS-S-002\s*费用明细", cell) for cell in source_cells
            ):
                continue
            tags.add(
                (
                    str(values[header["一级标签名"]] or "").strip(),
                    str(values[header["二级标签名"]] or "").strip(),
                )
            )
        return tags
    finally:
        workbook.close()


def load_medical_record_sources(
    path: Path,
) -> dict[tuple[str, str], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["拆分映射"]
        header: dict[str, int] | None = None
        sources: dict[tuple[str, str], list[str]] = {}
        for values in worksheet.iter_rows(values_only=True):
            if values and values[0] == "mapping_id":
                header = {
                    str(value): index
                    for index, value in enumerate(values)
                    if value is not None
                }
                continue
            if header is None or not values or not values[0]:
                continue
            record_type = str(values[header["recordType"]] or "").strip()
            if record_type != "病程记录":
                continue
            source_level = str(values[header["来源级别"]] or "").strip()
            source_column = (
                "主要来源"
                if source_level in {"PRIMARY", "主要来源"}
                else "次要来源"
                if source_level in {"SECONDARY", "次要来源"}
                else ""
            )
            if not source_column:
                raise ValueError(f"无法识别来源级别：{source_level}")
            source = str(values[header[source_column]] or "").strip()
            if not source or source == "—":
                continue
            key = (
                str(values[header["一级标签名"]] or "").strip(),
                str(values[header["二级标签名"]] or "").strip(),
            )
            tag_sources = sources.setdefault(key, [])
            if source not in tag_sources:
                tag_sources.append(source)
        return sources
    finally:
        workbook.close()

